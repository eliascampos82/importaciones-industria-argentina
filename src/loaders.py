"""Funciones para extraer las series por rama de los archivos crudos de INDEC / OEDE / ARCA.

Cada boletín tiene un layout distinto (headers de trimestre con texto inconsistente
entre años, headers de fecha con día variable, headers de año simple, bloques
"Año XXXX" con 12 filas de mes) - una función por formato. Todas devuelven data
"tidy" con la columna `rama` de src/ramas.py.
"""
import re
from pathlib import Path

import openpyxl
import pandas as pd
import xlrd

from ramas import RAMAS, TOTAL, ciiu2_to_rama, ncm_to_rama

RAW = Path(__file__).resolve().parents[1] / "data" / "raw"

_QUARTER_RE = re.compile(r"^\s*(\d)\D*[Tt]rim\D*?(\d{4})\s*$")
_MESES = {
    "Enero": 1, "Febrero": 2, "Marzo": 3, "Abril": 4, "Mayo": 5, "Junio": 6,
    "Julio": 7, "Agosto": 8, "Septiembre": 9, "Octubre": 10, "Noviembre": 11, "Diciembre": 12,
}


def _limpiar(texto) -> str:
    return " ".join(str(texto).split())


# --------------------------------------------------------------------------- empleo
def load_empleo_trimestral() -> pd.DataFrame:
    """Empleo asalariado privado registrado en manufactura, por CIIU2, trimestral 1996-2026."""
    wb = openpyxl.load_workbook(RAW / "empleo" / "nacional_serie_empleo_trimestral_8.xlsx",
                                read_only=True, data_only=True)
    rows = list(wb["C4"].iter_rows(values_only=True))
    header = rows[2]

    quarter_cols = []
    for i, h in enumerate(header):
        if isinstance(h, str) and "/" not in h:   # las columnas con "/" son variaciones, no datos
            m = _QUARTER_RE.match(h)
            if m:
                quarter_cols.append((i, int(m.group(2)), int(m.group(1))))

    data = []
    for r in rows[3:]:
        try:
            code = int(r[0])
        except (TypeError, ValueError):
            continue   # filas de subtotal sin código
        for i, y, q in quarter_cols:
            if r[i] is not None:
                data.append((code, y, q, r[i]))

    df = pd.DataFrame(data, columns=["ciiu2", "anio", "trimestre", "puestos"])
    df["periodo"] = pd.PeriodIndex.from_fields(year=df.anio, quarter=df.trimestre, freq="Q").to_timestamp()
    df["rama"] = df["ciiu2"].map(ciiu2_to_rama)
    return df.dropna(subset=["rama"]).reset_index(drop=True)


def load_empleo_industria_oficial() -> pd.DataFrame:
    """Fila 'Industria' (subtotal oficial) del cuadro 4, para verificar la suma de ramas."""
    wb = openpyxl.load_workbook(RAW / "empleo" / "nacional_serie_empleo_trimestral_8.xlsx",
                                read_only=True, data_only=True)
    rows = list(wb["C4"].iter_rows(values_only=True))
    header = rows[2]
    quarter_cols = [(i, int(_QUARTER_RE.match(h).group(2)), int(_QUARTER_RE.match(h).group(1)))
                    for i, h in enumerate(header)
                    if isinstance(h, str) and "/" not in h and _QUARTER_RE.match(h)]
    fila = next(r for r in rows[3:] if r[0] is None and isinstance(r[1], str) and r[1].strip() == "Industria")
    df = pd.DataFrame([(y, q, fila[i]) for i, y, q in quarter_cols], columns=["anio", "trimestre", "puestos_oficial"])
    df["periodo"] = pd.PeriodIndex.from_fields(year=df.anio, quarter=df.trimestre, freq="Q").to_timestamp()
    return df[["periodo", "puestos_oficial"]]


# ------------------------------------------------------------------------- empresas
def load_empresas_anual() -> pd.DataFrame:
    """Empresas privadas activas en manufactura, por CIIU2, anual 1996-2024."""
    wb = openpyxl.load_workbook(RAW / "empresas" / "nacional_serie_empresas_2.xlsx",
                                read_only=True, data_only=True)
    rows = list(wb["C 3"].iter_rows(values_only=True))
    year_cols = [(i, int(h)) for i, h in enumerate(rows[4]) if isinstance(h, (int, float))]

    data = []
    for r in rows[5:]:
        try:
            code = int(r[0])
        except (TypeError, ValueError):
            continue
        for i, y in year_cols:
            if r[i] is not None:
                data.append((code, y, r[i]))

    df = pd.DataFrame(data, columns=["ciiu2", "anio", "empresas"])
    df["rama"] = df["ciiu2"].map(ciiu2_to_rama)
    return df.dropna(subset=["rama"]).reset_index(drop=True)


# ------------------------------------------------------------------------- salarios
def load_remuneraciones_mensual() -> pd.DataFrame:
    """Remuneración promedio nominal por todo concepto (incluye aguinaldo), por CIIU2, mensual 1995-2026."""
    wb = openpyxl.load_workbook(RAW / "salarios" / "nacional_serie_remuneraciones_mensual_8.xlsx",
                                read_only=True, data_only=True)
    rows = list(wb["C 7"].iter_rows(values_only=True))
    date_cols = [(i, h) for i, h in enumerate(rows[4]) if hasattr(h, "year")]

    data = []
    for r in rows[5:]:
        try:
            code = int(r[0])
        except (TypeError, ValueError):
            continue
        for i, dt in date_cols:
            if r[i] is not None:
                data.append((code, dt.year, dt.month, r[i]))

    df = pd.DataFrame(data, columns=["ciiu2", "anio", "mes", "salario_nominal"])
    # el día del header es inconsistente en el Excel original: se reconstruye solo con año/mes
    df["periodo"] = pd.to_datetime(dict(year=df.anio, month=df.mes, day=1))
    df["rama"] = df["ciiu2"].map(ciiu2_to_rama)
    return df.dropna(subset=["rama"]).drop(columns=["anio", "mes"]).reset_index(drop=True)


def load_remuneracion_industria_oficial() -> pd.DataFrame:
    """Remuneración promedio nominal de 'Industrias manufactureras' (letra D), para verificar el promedio ponderado."""
    wb = openpyxl.load_workbook(RAW / "salarios" / "nacional_serie_remuneraciones_mensual_8.xlsx",
                                read_only=True, data_only=True)
    rows = list(wb["C 5.1"].iter_rows(values_only=True))
    col_d = list(rows[4]).index("D")
    data = [(dt.year, dt.month, r[col_d]) for r in rows[6:]
            if hasattr((dt := r[0]), "year") and r[col_d] is not None]
    df = pd.DataFrame(data, columns=["anio", "mes", "salario_oficial"])
    df["periodo"] = pd.to_datetime(dict(year=df.anio, month=df.mes, day=1))
    return df[["periodo", "salario_oficial"]]


def load_ipc_nacional() -> pd.DataFrame:
    """IPC nivel general nacional, mensual desde dic-2016 (antes no hay serie nacional continua)."""
    df = pd.read_csv(RAW / "precios" / "serie_ipc_divisiones.csv", sep=";", encoding="latin1")
    df = df[(df["Region"] == "Nacional") & (df["Descripcion"] == "NIVEL GENERAL")].copy()
    df["periodo"] = pd.to_datetime(df["Periodo"].astype(str), format="%Y%m")
    df["ipc"] = df["Indice_IPC"].astype(str).str.replace(",", ".").astype(float)
    return df[["periodo", "ipc"]].sort_values("periodo").reset_index(drop=True)


# ------------------------------------------------------------------------ industria
def load_capacidad_instalada() -> pd.DataFrame:
    """Utilización de la capacidad instalada (%), nivel general + bloques sectoriales, mensual 2016-2026."""
    ws = xlrd.open_workbook(RAW / "industria" / "capacidad_instalada_2016_2026.xls").sheet_by_name("UCI - NG y bloques")
    header = [_limpiar(h) for h in ws.row_values(2)]

    data, year = [], None
    for r in range(3, ws.nrows):
        row = ws.row_values(r)
        first = _limpiar(row[0])
        m = re.match(r"A.o (\d{4})", first)
        if m:
            year = int(m.group(1))
            continue
        mes = first.rstrip("*").strip()
        if mes in _MESES and year is not None:
            for c in range(1, len(header)):
                if isinstance(row[c], (int, float)):
                    data.append((year, _MESES[mes], header[c], row[c]))

    df = pd.DataFrame(data, columns=["anio", "mes", "bloque", "uci"])
    df["periodo"] = pd.to_datetime(dict(year=df.anio, month=df.mes, day=1))
    return df[["periodo", "bloque", "uci"]]


def load_ipi_manufacturero() -> pd.DataFrame:
    """IPI manufacturero (base 2004=100), mensual 2016-2026: nivel general, 16 ramas y subramas.

    nivel: 'total' (nivel general), 'rama' (una de las 16 agrupaciones), 'subrama' (detalle,
    ej. calzado, muebles); `rama` de una subrama es la rama que la contiene.
    """
    ws = xlrd.open_workbook(RAW / "industria" / "ipi_manufacturero_serie_2016_2026.xls").sheet_by_name("Cuadro 2")
    codes, names = ws.row_values(2), ws.row_values(3)

    cols = []   # (col, codigo, nombre, nivel, rama)
    for i, (c, n) in enumerate(zip(codes, names)):
        if i < 3 or not _limpiar(n):
            continue
        if isinstance(c, str) and c.strip() == "Nivel general":
            cols.append((i, TOTAL, _limpiar(n), "total", TOTAL))
            continue
        codigo = str(int(c)) if isinstance(c, float) and c == int(c) else _limpiar(c)
        if codigo in RAMAS:
            cols.append((i, codigo, _limpiar(n), "rama", codigo))
        else:
            cols.append((i, codigo, _limpiar(n), "subrama", ciiu2_to_rama(int(re.match(r"\d{2}", codigo).group()))))

    data, year = [], None
    for r in range(6, ws.nrows):
        row = ws.row_values(r)
        y = row[1]
        if isinstance(y, (int, float)) and y:
            year = int(y)
        elif isinstance(y, str) and re.match(r"^\d{4}", y.strip()):
            year = int(y.strip()[:4])      # los años provisorios vienen como '2023*'
        mes = str(row[2]).strip()
        if mes in _MESES and year is not None:
            for i, codigo, nombre, nivel, rama in cols:
                if isinstance(row[i], (int, float)):
                    data.append((year, _MESES[mes], codigo, nombre, nivel, rama, row[i]))

    df = pd.DataFrame(data, columns=["anio", "mes", "codigo", "nombre", "nivel", "rama", "ipi"])
    df["periodo"] = pd.to_datetime(dict(year=df.anio, month=df.mes, day=1))
    return df.drop(columns=["anio", "mes"])


# ------------------------------------------------------------------ comercio exterior
def load_paises() -> pd.DataFrame:
    """Tabla oficial de códigos de país (ARCA, formulario F.2668)."""
    ws = xlrd.open_workbook(RAW / "comercio_exterior" / "Tablas-Paises-F2668.xls").sheet_by_name("E_PAIS_V2")
    rows = [ws.row_values(r)[0].split(",", 1) for r in range(1, ws.nrows)]
    df = pd.DataFrame(rows, columns=["porg", "pais"])
    df["porg"] = df["porg"].astype(int)
    df["pais"] = df["pais"].str.strip()
    return df


def _numero(serie: pd.Series) -> pd.Series:
    """Texto con coma decimal y espacios -> float. Los valores confidenciales ('s', 's1'...) quedan NaN."""
    return pd.to_numeric(serie.str.strip().str.replace(",", ".", regex=False), errors="coerce")


def _leer_comex(ruta: Path, columnas: list[str]) -> pd.DataFrame:
    df = pd.read_csv(ruta, sep=";", dtype=str, encoding="latin1")
    df.columns = columnas
    df["ncm"] = df["ncm"].str.strip().str.zfill(8)
    for c in columnas[3:]:
        df[c] = _numero(df[c])
    df["anio"] = df["anio"].astype(int)
    df[columnas[2]] = df[columnas[2]].astype(int)
    return df


def load_importaciones(anio: int) -> pd.DataFrame:
    """Importaciones por NCM y país de origen de un año, con rama y nombre de país.

    2016-2018 publican solo CIF; desde 2019 también FOB, flete y seguro (fob_usd queda NaN antes).
    """
    aa = str(anio)[-2:]
    ruta = next((RAW / "comercio_exterior" / "importaciones").glob(f"impo[at]{aa}.csv"))
    ncols = len(pd.read_csv(ruta, sep=";", nrows=1, encoding="latin1").columns)
    columnas = (["anio", "ncm", "porg", "peso_neto_kg", "cif_usd"] if ncols == 5 else
                ["anio", "ncm", "porg", "peso_neto_kg", "fob_usd", "flete_usd", "seguro_usd", "cif_usd"])
    df = _leer_comex(ruta, columnas)
    if "fob_usd" not in df:
        df["fob_usd"] = float("nan")
    df["rama"] = df["ncm"].map(ncm_to_rama)
    df = df.merge(load_paises().rename(columns={"porg": "porg"}), on="porg", how="left")
    df["pais"] = df["pais"].fillna("No determinado / otros")
    return df


def control_importaciones(anio: int) -> dict:
    """Totales de control que publica INDEC para el archivo de importaciones de ese año."""
    aa = str(anio)[-2:]
    ruta = next((RAW / "comercio_exterior" / "importaciones").glob(f"itot[at]{aa}.csv"))
    c = pd.read_csv(ruta, sep=";", dtype=str, encoding="latin1")
    c.columns = [x.strip() for x in c.columns]
    fila = c.iloc[0]
    cif = next(col for col in c.columns if col.startswith("Total CIF"))
    return {"registros": int(_numero(pd.Series([fila["Total de registros"]])).iloc[0]),
            "cif_usd": float(_numero(pd.Series([fila[cif]])).iloc[0])}


def load_exportaciones(anio: int) -> pd.DataFrame:
    """Exportaciones por NCM y país de destino de un año (tabla que prioriza el producto).

    Los valores confidenciales ('s1'...'s10') quedan NaN y el NCM '99999999' (confidencial) sin rama.
    """
    aa = str(anio)[-2:]
    ruta = next((RAW / "comercio_exterior" / "exportaciones").glob(f"expo[tn]*{aa}.csv"))
    df = _leer_comex(ruta, ["anio", "ncm", "pdes", "peso_neto_kg", "fob_usd"])
    df["rama"] = df["ncm"].map(ncm_to_rama)
    return df


def control_exportaciones(anio: int) -> dict:
    aa = str(anio)[-2:]
    ruta = next((RAW / "comercio_exterior" / "exportaciones").glob(f"etot[tn]*{aa}.csv"))
    c = pd.read_csv(ruta, sep=";", dtype=str, encoding="latin1")
    c.columns = [x.strip() for x in c.columns]
    fila = c.iloc[0]
    fob = next(col for col in c.columns if col.startswith("Total FOB"))
    return {"registros": int(_numero(pd.Series([fila["Total de registros"]])).iloc[0]),
            "fob_usd": float(_numero(pd.Series([fila[fob]])).iloc[0])}
